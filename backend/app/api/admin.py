"""Admin management API routes.

Follows `docs/05-api-specification.md` (Sections 12, 14, 15), `docs/06-security-specification.md`,
and `docs/13-final-implementation-blueprint.md` (Phase 4).
"""

from __future__ import annotations

import io
import uuid
from datetime import date, datetime, timedelta, timezone
from typing import Annotated, Any

from fastapi import APIRouter, Depends, File, Query, Response, UploadFile, status
from fastapi.concurrency import run_in_threadpool
from sqlalchemy import desc, func, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.database import get_db
from app.core.dependencies import require_roles
from app.core.exceptions import DomainValidationError, ResourceNotFoundError
from app.models.billing import Bill, Payment
from app.models.enums import UserRole
from app.models.identity import User
from app.models.operations import AuditLog
from app.models.qr_sessions import TableSession
from app.schemas.audit import AuditLogFacets, AuditLogPageResponse, AuditLogResponse
from app.schemas.billing import BillResponse, PaymentResponse
from app.schemas.common import PaginationMeta
from app.schemas.menu import (
    CategoryReorderRequest,
    ItemReorderRequest,
    MenuCategoryCreate,
    MenuCategoryResponse,
    MenuCategoryUpdate,
    MenuItemAvailabilityUpdate,
    MenuItemCreate,
    MenuItemResponse,
    MenuItemUpdate,
    MenuPriceVariantCreate,
    MenuPriceVariantResponse,
    MenuPriceVariantUpdate,
)
from app.schemas.order import (
    OrderCancelRequest,
    SessionLimitOverrideRequest,
    WaiterOrderResponse,
)
from app.schemas.qr import (
    QrPrintCardResponse,
    QrRegenerateRequest,
    QrRegenerateResponse,
)
from app.schemas.reports import (
    AdminDashboardKPIs,
    CancelledRejectedReportItem,
    DailySalesSummary,
    KitchenActivityReportItem,
    OrderReportItem,
    OrderStatusReportItem,
    PaymentMethodReportItem,
    PopularDishReportItem,
    RevenueTimeSeriesItem,
    TableUsageReportItem,
    UnavailableDishReportItem,
    WaiterActivityReportItem,
)
from app.schemas.settings import (
    RestaurantSettingsResponse,
    RestaurantSettingsUpdate,
)
from app.schemas.staff import (
    StaffPasswordResetRequest,
    StaffUserAdminResponse,
    StaffUserCreate,
    StaffUserUpdate,
)
from app.schemas.table import (
    TableCreate,
    TableResponse,
    TableUpdate,
)
from app.services.audit_service import AuditService
from app.services.image_service import ImageService
from app.services.menu_service import MenuService
from app.services.order_service import OrderService
from app.services.qr_service import QrService
from app.services.report_service import ReportService
from app.services.settings_service import SettingsService
from app.services.staff_service import StaffService
from app.services.table_service import TableService

router = APIRouter(prefix="/admin", tags=["Admin Operations"])

AdminUser = Annotated[User, Depends(require_roles(UserRole.ADMIN))]
AdminOrKitchenUser = Annotated[User, Depends(require_roles(UserRole.ADMIN, UserRole.KITCHEN))]



# =============================================================================
# Menu Category Management (Section 12.1)
# =============================================================================


@router.get("/menu/categories", response_model=list[MenuCategoryResponse])
async def list_menu_categories(
    _admin: AdminUser,
    active_only: bool = Query(default=False),
    db: AsyncSession = Depends(get_db),
) -> list[MenuCategoryResponse]:
    """List all menu categories with display ordering."""
    categories = await MenuService.list_categories(db, active_only=active_only)
    return [MenuCategoryResponse.model_validate(c) for c in categories]


@router.post(
    "/menu/categories",
    response_model=MenuCategoryResponse,
    status_code=status.HTTP_201_CREATED,
)
async def create_menu_category(
    payload: MenuCategoryCreate,
    admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> MenuCategoryResponse:
    """Create a new menu category with unique name validation."""
    category = await MenuService.create_category(
        db,
        payload,
        actor_user_id=admin.id,
        actor_role=admin.role,
    )
    await db.commit()
    return MenuCategoryResponse.model_validate(category)


@router.put("/menu/categories/reorder", status_code=status.HTTP_200_OK)
async def reorder_menu_categories(
    payload: CategoryReorderRequest,
    admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> dict[str, str]:
    """Bulk update display orders for categories."""
    await MenuService.reorder_categories(
        db,
        payload,
        actor_user_id=admin.id,
        actor_role=admin.role,
    )
    await db.commit()
    return {"status": "ok"}


@router.patch("/menu/categories/{category_id}", response_model=MenuCategoryResponse)
async def update_menu_category(
    category_id: uuid.UUID,
    payload: MenuCategoryUpdate,
    admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> MenuCategoryResponse:
    """Update category properties or soft-deactivate."""
    category = await MenuService.update_category(
        db,
        category_id,
        payload,
        actor_user_id=admin.id,
        actor_role=admin.role,
    )
    await db.commit()
    return MenuCategoryResponse.model_validate(category)


@router.delete("/menu/categories/{category_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_menu_category(
    category_id: uuid.UUID,
    admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> Response:
    """Soft-deactivate category safely preserving historical orders."""
    await MenuService.delete_category(
        db,
        category_id,
        actor_user_id=admin.id,
        actor_role=admin.role,
    )
    await db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


# =============================================================================
# Menu Item & Short Code Management (Section 12.2)
# =============================================================================


@router.get("/menu/items", response_model=list[MenuItemResponse])
async def list_menu_items(
    _admin: AdminUser,
    category_id: uuid.UUID | None = Query(default=None),
    active_only: bool = Query(default=False),
    db: AsyncSession = Depends(get_db),
) -> list[MenuItemResponse]:
    """List dishes with their nested variants."""
    items = await MenuService.list_items(db, category_id=category_id, active_only=active_only)
    return [MenuItemResponse.model_validate(i) for i in items]


@router.post(
    "/menu/items",
    response_model=MenuItemResponse,
    status_code=status.HTTP_201_CREATED,
)
async def create_menu_item(
    payload: MenuItemCreate,
    admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> MenuItemResponse:
    """Create a new dish along with its initial price variant(s) in one atomic transaction."""
    item = await MenuService.create_item(
        db,
        payload,
        actor_user_id=admin.id,
        actor_role=admin.role,
    )
    await db.commit()
    return MenuItemResponse.model_validate(item)


@router.put("/menu/items/reorder", status_code=status.HTTP_200_OK)
async def reorder_menu_items(
    payload: ItemReorderRequest,
    admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> dict[str, str]:
    """Bulk update display orders for dishes within a category."""
    await MenuService.reorder_items(
        db,
        payload,
        actor_user_id=admin.id,
        actor_role=admin.role,
    )
    await db.commit()
    return {"status": "ok"}


@router.patch("/menu/items/{item_id}", response_model=MenuItemResponse)
async def update_menu_item(
    item_id: uuid.UUID,
    payload: MenuItemUpdate,
    admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> MenuItemResponse:
    """Update dish details or soft-deactivate."""
    item = await MenuService.update_item(
        db,
        item_id,
        payload,
        actor_user_id=admin.id,
        actor_role=admin.role,
    )
    await db.commit()
    return MenuItemResponse.model_validate(item)


@router.patch("/menu/items/{item_id}/availability", response_model=MenuItemResponse)
async def set_menu_item_availability(
    item_id: uuid.UUID,
    payload: MenuItemAvailabilityUpdate,
    staff: AdminOrKitchenUser,
    db: AsyncSession = Depends(get_db),
) -> MenuItemResponse:
    """Instant stock availability toggle for Admin and Kitchen."""
    # OD-05: same switch the KDS endpoint honours, so Kitchen cannot bypass it via this route.
    if staff.role != UserRole.ADMIN.value:
        cfg = await SettingsService.get_settings(db)
        if not cfg.kitchen_can_toggle_availability:
            raise DomainValidationError(
                "KITCHEN_TOGGLE_DISABLED",
                "Kitchen staff is not permitted to toggle item availability. Please ask an Admin.",
            )
    item = await MenuService.set_item_availability(
        db,
        item_id,
        is_available=payload.is_available,
        reason=payload.reason,
        actor_user_id=staff.id,
        actor_role=staff.role,
    )
    await db.commit()
    return MenuItemResponse.model_validate(item)


# =============================================================================
# Price Variant Management (Section 12.3)
# =============================================================================


@router.post(
    "/menu/items/{item_id}/variants",
    response_model=MenuPriceVariantResponse,
    status_code=status.HTTP_201_CREATED,
)
async def create_menu_variant(
    item_id: uuid.UUID,
    payload: MenuPriceVariantCreate,
    admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> MenuPriceVariantResponse:
    """Add a new portion/price variant to an existing dish."""
    variant = await MenuService.create_variant(
        db,
        item_id,
        payload,
        actor_user_id=admin.id,
        actor_role=admin.role,
    )
    await db.commit()
    return MenuPriceVariantResponse.model_validate(variant)


@router.patch("/menu/variants/{variant_id}", response_model=MenuPriceVariantResponse)
async def update_menu_variant(
    variant_id: uuid.UUID,
    payload: MenuPriceVariantUpdate,
    admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> MenuPriceVariantResponse:
    """Update a price variant (historical prices on existing orders remain immutable)."""
    variant = await MenuService.update_variant(
        db,
        variant_id,
        payload,
        actor_user_id=admin.id,
        actor_role=admin.role,
    )
    await db.commit()
    return MenuPriceVariantResponse.model_validate(variant)


# =============================================================================
# Image Upload & Sanitization (Section 26)
# =============================================================================


@router.post("/menu/items/{item_id}/image", response_model=MenuItemResponse)
async def upload_menu_item_image(
    item_id: uuid.UUID,
    admin: AdminUser,
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
) -> MenuItemResponse:
    """Upload dish photo. Sanitized via Pillow, EXIF stripped, converted to WebP."""
    item = await MenuService.get_item(db, item_id)

    raw_bytes = await file.read()
    if not raw_bytes:
        raise DomainValidationError(
            code="EMPTY_FILE",
            message="Uploaded file has no content.",
        )

    # Sanitize, strip EXIF, convert to WebP, and save under random UUID filename
    image_url, _filename = ImageService.sanitize_and_save_menu_image(raw_bytes)

    # Clean up prior image if replacing
    if item.image_url and item.image_url != image_url:
        ImageService.delete_menu_image(item.image_url)

    item.image_url = image_url
    await db.flush()

    await AuditService.record_event(
        db,
        action_type="MENU_ITEM_IMAGE_UPLOADED",
        entity_type="menu_item",
        entity_id=str(item.id),
        actor_role=admin.role,
        actor_user_id=admin.id,
        new_value={"image_url": image_url},
    )

    await db.commit()
    return MenuItemResponse.model_validate(item)


@router.delete(
    "/menu/items/{item_id}/image",
    status_code=status.HTTP_204_NO_CONTENT,
    response_class=Response,
)
async def delete_menu_item_image(
    item_id: uuid.UUID,
    admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> Response:
    """Remove image from a dish and delete file from disk."""
    item = await MenuService.get_item(db, item_id)
    if item.image_url:
        old_url = item.image_url
        ImageService.delete_menu_image(old_url)
        item.image_url = None
        await db.flush()

        await AuditService.record_event(
            db,
            action_type="MENU_ITEM_IMAGE_DELETED",
            entity_type="menu_item",
            entity_id=str(item.id),
            actor_role=admin.role,
            actor_user_id=admin.id,
            old_value={"image_url": old_url},
        )
        await db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


# =============================================================================
# Physical Table & QR Management (Section 12.3 & 12.4)
# =============================================================================


@router.get("/tables", response_model=list[TableResponse])
async def list_admin_tables(
    _admin: AdminUser,
    active_only: bool = Query(default=False),
    floor_or_zone: str | None = Query(default=None),
    db: AsyncSession = Depends(get_db),
) -> list[TableResponse]:
    """List physical dining tables with derived occupancy state and active QR info."""
    return await TableService.list_tables(db, active_only=active_only, floor_or_zone=floor_or_zone)


@router.post("/tables", response_model=TableResponse, status_code=status.HTTP_201_CREATED)
async def create_table(
    payload: TableCreate,
    admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> TableResponse:
    """Create a new physical dining table and automatically issue its initial v1 QR token."""
    res = await TableService.create_table(db, payload, user_id=admin.id)
    await db.commit()
    return res


@router.get("/tables/{table_id}", response_model=TableResponse)
async def get_table_details(
    table_id: uuid.UUID,
    _admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> TableResponse:
    """Retrieve details and occupancy state for a single table."""
    return await TableService.get_table(db, table_id)


@router.patch("/tables/{table_id}", response_model=TableResponse)
async def update_table(
    table_id: uuid.UUID,
    payload: TableUpdate,
    admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> TableResponse:
    """Update physical table parameters. Deactivating requires no active dining session."""
    res = await TableService.update_table(db, table_id, payload, user_id=admin.id)
    await db.commit()
    return res


@router.delete("/tables/{table_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_table(
    table_id: uuid.UUID,
    admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> Response:
    """Delete a never-used table. Tables with order history must be deactivated instead."""
    await TableService.delete_table(db, table_id, user_id=admin.id)
    await db.commit()
    return Response(status_code=status.HTTP_204_NO_CONTENT)


@router.post("/tables/{table_id}/qr/regenerate", response_model=QrRegenerateResponse)
async def regenerate_table_qr(
    table_id: uuid.UUID,
    payload: QrRegenerateRequest,
    admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> QrRegenerateResponse:
    """Atomically rotate a table's QR token, revoking the previous version."""
    res = await QrService.regenerate_qr_token(
        db, table_id, user_id=admin.id, reason=payload.reason or "ROTATED_BY_ADMIN"
    )
    await db.commit()
    return res


@router.get("/tables/{table_id}/qr-print", response_model=QrPrintCardResponse)
async def get_printable_qr_card(
    table_id: uuid.UUID,
    _admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> QrPrintCardResponse:
    """Retrieve decrypted printable standee card details for Admin."""
    return await QrService.get_printable_qr_card(db, table_id)


# =============================================================================
# Dashboard Overview KPIs (Section 11.1)
# =============================================================================


@router.get("/dashboard", response_model=AdminDashboardKPIs)
async def get_admin_dashboard(
    _admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> AdminDashboardKPIs:
    """Retrieve live real-time floor KPIs for the Admin Overview screen."""
    return await ReportService.get_dashboard_kpis(db)


# =============================================================================
# Staff Account Management (Section 11.7)
# =============================================================================


@router.get("/users", response_model=list[StaffUserAdminResponse])
async def list_staff_users(
    _admin: AdminUser,
    role: UserRole | None = Query(default=None),
    active_only: bool = Query(default=False),
    db: AsyncSession = Depends(get_db),
) -> list[StaffUserAdminResponse]:
    """List staff accounts with role and active filters."""
    return await StaffService.list_users(db, role=role, active_only=active_only)


@router.post(
    "/users",
    response_model=StaffUserAdminResponse,
    status_code=status.HTTP_201_CREATED,
)
async def create_staff_user(
    payload: StaffUserCreate,
    admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> StaffUserAdminResponse:
    """Create a new WAITER, KITCHEN, or ADMIN account."""
    res = await StaffService.create_user(db, payload, admin_user=admin)
    await db.commit()
    return res


@router.patch("/users/{user_id}", response_model=StaffUserAdminResponse)
async def update_staff_user(
    user_id: uuid.UUID,
    payload: StaffUserUpdate,
    admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> StaffUserAdminResponse:
    """Update staff details, role, or immediately deactivate."""
    res = await StaffService.update_user(db, user_id, payload, admin_user=admin)
    await db.commit()
    return res


@router.post("/users/{user_id}/reset-password", response_model=StaffUserAdminResponse)
async def reset_staff_password(
    user_id: uuid.UUID,
    payload: StaffPasswordResetRequest,
    admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> StaffUserAdminResponse:
    """Reset staff password, invalidate existing tokens, and disconnect sessions."""
    res = await StaffService.reset_password(db, user_id, payload, admin_user=admin)
    await db.commit()
    return res


@router.post("/users/{user_id}/unlock", response_model=StaffUserAdminResponse)
async def unlock_staff_user(
    user_id: uuid.UUID,
    admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> StaffUserAdminResponse:
    """Clear failed login attempts and unlock account."""
    res = await StaffService.unlock_user(db, user_id, admin_user=admin)
    await db.commit()
    return res


# =============================================================================
# Restaurant Settings Management (Section 11.10)
# =============================================================================


@router.get("/settings", response_model=RestaurantSettingsResponse)
async def get_restaurant_settings(
    _admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> RestaurantSettingsResponse:
    """Retrieve singleton restaurant configuration and policy switches."""
    return await SettingsService.get_settings_response(db)


@router.patch("/settings", response_model=RestaurantSettingsResponse)
async def update_restaurant_settings(
    payload: RestaurantSettingsUpdate,
    admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> RestaurantSettingsResponse:
    """Update restaurant configuration with audit logging."""
    res = await SettingsService.update_settings(db, payload, admin_user=admin)
    await db.commit()
    return res


# =============================================================================
# Audit Log Viewer (Section 23)
# =============================================================================


# India has no DST, so a fixed +05:30 offset is exact for restaurant-day boundaries.
_IST = timezone(timedelta(hours=5, minutes=30))
_AUDIT_EXPORT_MAX_ROWS = 50_000


def _escape_like(value: str) -> str:
    return value.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")


def _audit_log_query(
    *,
    action_type: str | None,
    entity_type: str | None,
    entity_id: str | None,
    actor_user_id: uuid.UUID | None,
    actor_role: str | None,
    from_date: date | None,
    to_date: date | None,
    q: str | None,
    columns: Any,
) -> Any:
    """Build a filtered audit query. Date bounds are IST calendar days (inclusive)."""
    stmt = select(*columns).select_from(AuditLog).outerjoin(User, User.id == AuditLog.actor_user_id)

    if action_type:
        stmt = stmt.where(AuditLog.action_type == action_type)
    if entity_type:
        stmt = stmt.where(AuditLog.entity_type == entity_type)
    if entity_id:
        stmt = stmt.where(AuditLog.entity_id == entity_id)
    if actor_user_id:
        stmt = stmt.where(AuditLog.actor_user_id == actor_user_id)
    if actor_role:
        stmt = stmt.where(AuditLog.actor_role == actor_role)
    if from_date:
        stmt = stmt.where(
            AuditLog.created_at >= datetime.combine(from_date, datetime.min.time(), tzinfo=_IST)
        )
    if to_date:
        stmt = stmt.where(
            AuditLog.created_at
            < datetime.combine(to_date + timedelta(days=1), datetime.min.time(), tzinfo=_IST)
        )
    if q and q.strip():
        pattern = f"%{_escape_like(q.strip())}%"
        stmt = stmt.where(
            AuditLog.action_type.ilike(pattern, escape="\\")
            | AuditLog.entity_type.ilike(pattern, escape="\\")
            | AuditLog.entity_id.ilike(pattern, escape="\\")
            | AuditLog.reason.ilike(pattern, escape="\\")
            | User.username.ilike(pattern, escape="\\")
        )
    return stmt


def _audit_response(log: AuditLog, username: str | None) -> AuditLogResponse:
    item = AuditLogResponse.model_validate(log)
    item.actor_username = username
    return item


@router.get("/audit-logs/facets", response_model=AuditLogFacets)
async def get_audit_log_facets(
    _admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> AuditLogFacets:
    """Distinct action types, entity types and actor roles currently in the audit trail."""

    async def _distinct(column: Any) -> list[str]:
        rows = await db.scalars(select(column).distinct().order_by(column))
        return [r for r in rows.all() if r]

    return AuditLogFacets(
        action_types=await _distinct(AuditLog.action_type),
        entity_types=await _distinct(AuditLog.entity_type),
        actor_roles=await _distinct(AuditLog.actor_role),
    )


@router.get("/audit-logs", response_model=AuditLogPageResponse)
async def list_audit_logs(
    _admin: AdminUser,
    action_type: str | None = Query(default=None),
    entity_type: str | None = Query(default=None),
    entity_id: str | None = Query(default=None),
    actor_user_id: uuid.UUID | None = Query(default=None),
    actor_role: str | None = Query(default=None),
    q: str | None = Query(default=None, max_length=100),
    from_date: date | None = Query(default=None, alias="from"),
    to_date: date | None = Query(default=None, alias="to"),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=20, ge=1, le=100),
    db: AsyncSession = Depends(get_db),
) -> AuditLogPageResponse:
    """Query immutable system audit logs."""
    filters: dict[str, Any] = {
        "action_type": action_type,
        "entity_type": entity_type,
        "entity_id": entity_id,
        "actor_user_id": actor_user_id,
        "actor_role": actor_role,
        "from_date": from_date,
        "to_date": to_date,
        "q": q,
    }
    count_stmt = _audit_log_query(**filters, columns=[func.count(AuditLog.id)])
    total_items = (await db.scalar(count_stmt)) or 0
    total_pages = (total_items + page_size - 1) // page_size if total_items > 0 else 0

    stmt = (
        _audit_log_query(**filters, columns=[AuditLog, User.username])
        .order_by(desc(AuditLog.created_at))
        .offset((page - 1) * page_size)
        .limit(page_size)
    )
    rows = (await db.execute(stmt)).all()

    return AuditLogPageResponse(
        items=[_audit_response(log, username) for log, username in rows],
        pagination=PaginationMeta(
            page=page,
            page_size=page_size,
            total_items=total_items,
            total_pages=total_pages,
            has_next=page < total_pages,
            has_prev=page > 1,
        ),
    )


def _build_audit_workbook(
    rows: list[tuple[AuditLog, str | None]],
    *,
    filter_summary: list[tuple[str, str]],
    generated_by: str,
    truncated: bool,
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    def _ist(dt: datetime) -> str:
        return dt.astimezone(_IST).strftime("%d %b %Y, %I:%M:%S %p")

    def _json(value: Any) -> str:
        if value is None:
            return ""
        import json

        return json.dumps(value, ensure_ascii=False, default=str)[:32000]

    def _put(ws: Any, row: int, col: int, value: Any) -> Any:
        cell = ws.cell(row=row, column=col, value=value)
        # Force text so values like "=SUM(...)" in a reason are never evaluated as formulas.
        if isinstance(value, str):
            cell.data_type = "s"
        return cell

    wb = Workbook()
    ws = wb.active
    ws.title = "Audit Log"
    headers = [
        ("Timestamp (IST)", 24),
        ("Action", 30),
        ("Actor", 18),
        ("Role", 14),
        ("Entity Type", 20),
        ("Entity ID", 38),
        ("Reason", 44),
        ("Previous State", 48),
        ("New State", 48),
        ("Metadata", 36),
    ]
    head_fill = PatternFill("solid", fgColor="D95D2B")
    for idx, (title, width) in enumerate(headers, start=1):
        c = _put(ws, 1, idx, title)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 22

    for r, (log, username) in enumerate(rows, start=2):
        values = [
            _ist(log.created_at),
            log.action_type,
            username or "—",
            log.actor_role,
            log.entity_type,
            log.entity_id,
            log.reason or "",
            _json(log.old_value_json),
            _json(log.new_value_json),
            _json(log.metadata_json),
        ]
        for c_idx, v in enumerate(values, start=1):
            _put(ws, r, c_idx, v).alignment = Alignment(vertical="top")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 1)}"

    info = wb.create_sheet("Export Info")
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 60
    summary = [
        ("Generated at (IST)", _ist(datetime.now(timezone.utc))),
        ("Generated by", generated_by),
        ("Rows exported", str(len(rows))),
        *filter_summary,
    ]
    if truncated:
        summary.append(
            (
                "Note",
                f"Result capped at {_AUDIT_EXPORT_MAX_ROWS:,} newest rows. "
                "Narrow the date range to export older entries.",
            )
        )
    for i, (k, v) in enumerate(summary, start=1):
        _put(info, i, 1, k).font = Font(bold=True)
        _put(info, i, 2, v)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/audit-logs/export.xlsx")
async def export_audit_logs_xlsx(
    admin: AdminUser,
    action_type: str | None = Query(default=None),
    entity_type: str | None = Query(default=None),
    entity_id: str | None = Query(default=None),
    actor_user_id: uuid.UUID | None = Query(default=None),
    actor_role: str | None = Query(default=None),
    q: str | None = Query(default=None, max_length=100),
    from_date: date | None = Query(default=None, alias="from"),
    to_date: date | None = Query(default=None, alias="to"),
    db: AsyncSession = Depends(get_db),
) -> Response:
    """Download every audit entry matching the filters (not just one page) as an .xlsx file."""
    filters: dict[str, Any] = {
        "action_type": action_type,
        "entity_type": entity_type,
        "entity_id": entity_id,
        "actor_user_id": actor_user_id,
        "actor_role": actor_role,
        "from_date": from_date,
        "to_date": to_date,
        "q": q,
    }
    stmt = (
        _audit_log_query(**filters, columns=[AuditLog, User.username])
        .order_by(desc(AuditLog.created_at))
        .limit(_AUDIT_EXPORT_MAX_ROWS + 1)
    )
    fetched = (await db.execute(stmt)).all()
    truncated = len(fetched) > _AUDIT_EXPORT_MAX_ROWS
    rows = [(log, username) for log, username in fetched[:_AUDIT_EXPORT_MAX_ROWS]]

    filter_summary = [
        ("Filter: search", q or "—"),
        ("Filter: action", action_type or "All"),
        ("Filter: entity", entity_type or "All"),
        ("Filter: role", actor_role or "All"),
        ("Filter: from (IST)", from_date.isoformat() if from_date else "—"),
        ("Filter: to (IST)", to_date.isoformat() if to_date else "—"),
    ]
    content = await run_in_threadpool(
        _build_audit_workbook,
        rows,
        filter_summary=filter_summary,
        generated_by=admin.username,
        truncated=truncated,
    )

    await AuditService.record_event(
        db,
        action_type="AUDIT_LOG_EXPORTED",
        entity_type="audit_log",
        entity_id="export",
        actor_role=admin.role,
        actor_user_id=admin.id,
        metadata={
            "rows": len(rows),
            "truncated": truncated,
            "action_type": action_type,
            "entity_type": entity_type,
            "actor_role": actor_role,
            "from": from_date.isoformat() if from_date else None,
            "to": to_date.isoformat() if to_date else None,
            "q": q,
        },
    )
    await db.commit()

    stamp = datetime.now(_IST).strftime("%Y%m%d_%H%M")
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=audit_log_{stamp}.xlsx"},
    )



# =============================================================================
# Operational & Financial Reports (Section 22)
# =============================================================================


@router.get("/reports/daily-sales", response_model=list[DailySalesSummary])
async def get_daily_sales_report(
    _admin: AdminUser,
    from_date: date | None = Query(default=None, alias="from"),
    to_date: date | None = Query(default=None, alias="to"),
    db: AsyncSession = Depends(get_db),
) -> list[DailySalesSummary]:
    """Gross subtotal, discounts, CGST, SGST, net revenue, and average bill value."""
    return await ReportService.get_daily_sales(db, from_date=from_date, to_date=to_date)


@router.get("/reports/orders", response_model=list[OrderReportItem])
async def get_orders_report(
    _admin: AdminUser,
    from_date: date | None = Query(default=None, alias="from"),
    to_date: date | None = Query(default=None, alias="to"),
    status_filter: str | None = Query(default=None, alias="status"),
    source_filter: str | None = Query(default=None, alias="source"),
    db: AsyncSession = Depends(get_db),
) -> list[OrderReportItem]:
    """Order log with breakdown by source and status."""
    return await ReportService.get_orders_report(
        db,
        from_date=from_date,
        to_date=to_date,
        status_filter=status_filter,
        source_filter=source_filter,
    )


@router.get("/reports/popular-dishes", response_model=list[PopularDishReportItem])
async def get_popular_dishes_report(
    _admin: AdminUser,
    from_date: date | None = Query(default=None, alias="from"),
    to_date: date | None = Query(default=None, alias="to"),
    limit: int = Query(default=50, ge=1, le=200),
    db: AsyncSession = Depends(get_db),
) -> list[PopularDishReportItem]:
    """Top-selling menu items and price variants ranked by quantity billed."""
    return await ReportService.get_popular_dishes(
        db, from_date=from_date, to_date=to_date, limit=limit
    )


@router.get("/reports/revenue", response_model=list[RevenueTimeSeriesItem])
async def get_revenue_time_series_report(
    _admin: AdminUser,
    from_date: date | None = Query(default=None, alias="from"),
    to_date: date | None = Query(default=None, alias="to"),
    db: AsyncSession = Depends(get_db),
) -> list[RevenueTimeSeriesItem]:
    """Time-series revenue breakdown."""
    return await ReportService.get_revenue_time_series(db, from_date=from_date, to_date=to_date)


@router.get("/reports/payment-methods", response_model=list[PaymentMethodReportItem])
async def get_payment_methods_report(
    _admin: AdminUser,
    from_date: date | None = Query(default=None, alias="from"),
    to_date: date | None = Query(default=None, alias="to"),
    db: AsyncSession = Depends(get_db),
) -> list[PaymentMethodReportItem]:
    """Collection totals and transaction counts grouped by payment method."""
    return await ReportService.get_payment_methods(db, from_date=from_date, to_date=to_date)


@router.get("/reports/table-usage", response_model=list[TableUsageReportItem])
async def get_table_usage_report(
    _admin: AdminUser,
    from_date: date | None = Query(default=None, alias="from"),
    to_date: date | None = Query(default=None, alias="to"),
    db: AsyncSession = Depends(get_db),
) -> list[TableUsageReportItem]:
    """Table turnover, average session duration, and revenue per table."""
    return await ReportService.get_table_usage(db, from_date=from_date, to_date=to_date)


@router.get("/reports/order-status", response_model=list[OrderStatusReportItem])
async def get_order_status_report(
    _admin: AdminUser,
    from_date: date | None = Query(default=None, alias="from"),
    to_date: date | None = Query(default=None, alias="to"),
    db: AsyncSession = Depends(get_db),
) -> list[OrderStatusReportItem]:
    """Live and historical counts of orders across statuses."""
    return await ReportService.get_order_status(db, from_date=from_date, to_date=to_date)


@router.get("/reports/unavailable-dishes", response_model=list[UnavailableDishReportItem])
async def get_unavailable_dishes_report(
    _admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> list[UnavailableDishReportItem]:
    """List of currently unavailable / 86'd menu items."""
    return await ReportService.get_unavailable_dishes(db)


@router.get("/reports/cancelled-rejected", response_model=list[CancelledRejectedReportItem])
async def get_cancelled_rejected_report(
    _admin: AdminUser,
    from_date: date | None = Query(default=None, alias="from"),
    to_date: date | None = Query(default=None, alias="to"),
    db: AsyncSession = Depends(get_db),
) -> list[CancelledRejectedReportItem]:
    """Accountability report of all rejected and cancelled orders."""
    return await ReportService.get_cancelled_rejected(db, from_date=from_date, to_date=to_date)


@router.get("/reports/waiter-activity", response_model=list[WaiterActivityReportItem])
async def get_waiter_activity_report(
    _admin: AdminUser,
    from_date: date | None = Query(default=None, alias="from"),
    to_date: date | None = Query(default=None, alias="to"),
    db: AsyncSession = Depends(get_db),
) -> list[WaiterActivityReportItem]:
    """Per-waiter operational metrics."""
    return await ReportService.get_waiter_activity(db, from_date=from_date, to_date=to_date)


@router.get("/reports/kitchen-activity", response_model=list[KitchenActivityReportItem])
async def get_kitchen_activity_report(
    _admin: AdminUser,
    from_date: date | None = Query(default=None, alias="from"),
    to_date: date | None = Query(default=None, alias="to"),
    db: AsyncSession = Depends(get_db),
) -> list[KitchenActivityReportItem]:
    """Per-kitchen staff turnaround metrics."""
    return await ReportService.get_kitchen_activity(db, from_date=from_date, to_date=to_date)


@router.get("/reports/export.csv")
async def export_sales_csv(
    _admin: AdminUser,
    from_date: date | None = Query(default=None, alias="from"),
    to_date: date | None = Query(default=None, alias="to"),
    db: AsyncSession = Depends(get_db),
) -> Response:
    """Download RFC 4180 CSV export of sales for accounting."""
    csv_data = await ReportService.generate_sales_csv(db, from_date=from_date, to_date=to_date)
    f_str = from_date.isoformat() if from_date else "start"
    t_str = to_date.isoformat() if to_date else "today"
    filename = f"sales_report_{f_str}_{t_str}.csv"
    return Response(
        content=csv_data,
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


# =============================================================================
# Order & Session Overrides (Section 11.8 & 11.9)
# =============================================================================


@router.get("/orders", response_model=list[OrderReportItem])
async def list_admin_orders(
    _admin: AdminUser,
    status_filter: str | None = Query(default=None, alias="status"),
    source_filter: str | None = Query(default=None, alias="source"),
    from_date: date | None = Query(default=None, alias="from"),
    to_date: date | None = Query(default=None, alias="to"),
    db: AsyncSession = Depends(get_db),
) -> list[OrderReportItem]:
    """Admin view of all floor orders with filtering."""
    return await ReportService.get_orders_report(
        db,
        from_date=from_date,
        to_date=to_date,
        status_filter=status_filter,
        source_filter=source_filter,
    )


@router.post("/orders/{order_id}/void", response_model=WaiterOrderResponse)
async def admin_void_order(
    order_id: uuid.UUID,
    payload: OrderCancelRequest,
    admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> WaiterOrderResponse:
    """Perform emergency void on an order up to SERVED state with mandatory audit reason."""
    res = await OrderService.cancel_order(
        db,
        order_id=order_id,
        payload=payload,
        staff_user=admin,
    )
    await db.commit()
    return res


@router.post("/table-sessions/{session_id}/override-limit")
async def override_session_outstanding_limit(
    session_id: uuid.UUID,
    payload: SessionLimitOverrideRequest,
    admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    """Admin override to raise a table session's unpaid exposure limit."""
    sess = await db.get(TableSession, session_id)
    if not sess:
        raise ResourceNotFoundError("SESSION_NOT_FOUND", f"Table session {session_id} not found.")

    old_limit = sess.custom_outstanding_limit
    sess.custom_outstanding_limit = payload.new_limit
    sess.limit_override_by_user_id = admin.id
    sess.limit_override_reason = payload.reason

    await AuditService.record_event(
        db,
        action_type="SESSION_LIMIT_OVERRIDDEN",
        entity_type="TABLE_SESSION",
        entity_id=str(sess.id),
        actor_role=admin.role,
        actor_user_id=admin.id,
        reason=payload.reason,
        old_value={"custom_outstanding_limit": str(old_limit)},
        new_value={"custom_outstanding_limit": str(payload.new_limit)},
    )
    await db.commit()
    return {
        "table_session_id": str(sess.id),
        "custom_outstanding_limit": str(sess.custom_outstanding_limit),
        "limit_override_reason": sess.limit_override_reason,
    }


@router.post("/table-sessions/{session_id}/unlock")
async def unlock_table_session(
    session_id: uuid.UUID,
    admin: AdminUser,
    db: AsyncSession = Depends(get_db),
) -> dict[str, Any]:
    """Clear abuse lock on a rate-limited table session."""
    sess = await db.get(TableSession, session_id)
    if not sess:
        raise ResourceNotFoundError("SESSION_NOT_FOUND", f"Table session {session_id} not found.")

    sess.is_locked_for_abuse = False
    await AuditService.record_event(
        db,
        action_type="SESSION_ABUSE_LOCK_CLEARED",
        entity_type="TABLE_SESSION",
        entity_id=str(sess.id),
        actor_role=admin.role,
        actor_user_id=admin.id,
        reason="Admin cleared session abuse lock",
    )
    await db.commit()
    return {"table_session_id": str(sess.id), "is_locked_for_abuse": False}


@router.get("/bills", response_model=list[BillResponse])
async def list_admin_bills(
    _admin: AdminUser,
    status_filter: str | None = Query(default=None, alias="status"),
    from_date: date | None = Query(default=None, alias="from"),
    to_date: date | None = Query(default=None, alias="to"),
    db: AsyncSession = Depends(get_db),
) -> list[BillResponse]:
    """List all bills with optional status and date filters."""
    stmt = (
        select(Bill)
        .options(selectinload(Bill.items), selectinload(Bill.payments))
        .order_by(Bill.generated_at.desc())
    )
    if status_filter:
        stmt = stmt.where(Bill.status == status_filter)
    if from_date:
        start_dt = datetime.combine(from_date, datetime.min.time(), tzinfo=timezone.utc)
        stmt = stmt.where(Bill.generated_at >= start_dt)
    if to_date:
        end_dt = datetime.combine(to_date, datetime.max.time(), tzinfo=timezone.utc)
        stmt = stmt.where(Bill.generated_at <= end_dt)

    bills = (await db.scalars(stmt)).all()
    return [BillResponse.model_validate(b) for b in bills]


@router.get("/payments", response_model=list[PaymentResponse])
async def list_admin_payments(
    _admin: AdminUser,
    status_filter: str | None = Query(default=None, alias="status"),
    payment_method: str | None = Query(default=None),
    from_date: date | None = Query(default=None, alias="from"),
    to_date: date | None = Query(default=None, alias="to"),
    db: AsyncSession = Depends(get_db),
) -> list[PaymentResponse]:
    """List all payment records with optional filters."""
    stmt = select(Payment).order_by(Payment.created_at.desc())
    if status_filter:
        stmt = stmt.where(Payment.status == status_filter)
    if payment_method:
        stmt = stmt.where(Payment.payment_method == payment_method)
    if from_date:
        start_dt = datetime.combine(from_date, datetime.min.time(), tzinfo=timezone.utc)
        stmt = stmt.where(Payment.created_at >= start_dt)
    if to_date:
        end_dt = datetime.combine(to_date, datetime.max.time(), tzinfo=timezone.utc)
        stmt = stmt.where(Payment.created_at <= end_dt)

    payments = (await db.scalars(stmt)).all()
    return [PaymentResponse.model_validate(p) for p in payments]

